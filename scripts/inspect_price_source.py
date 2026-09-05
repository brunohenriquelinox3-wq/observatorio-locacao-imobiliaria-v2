"""Inspeção mínima e saneada de planilha para a prévia de importação A223.

Não imprime nem grava conteúdo de células. Produz apenas metadados, cabeçalhos,
contagens agregadas, sinais de fórmula/link e aderência de identificadores.
"""

from __future__ import annotations

import json
import re
import sys
from collections import Counter
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def normalize(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def classify(value: Any) -> str:
    if value is None or value == "":
        return "empty"
    if is_number(value):
        return "numeric"
    if hasattr(value, "isoformat"):
        return "date"
    return "text"


def index_for(headers: list[str], patterns: tuple[str, ...]) -> int | None:
    for index, header in enumerate(headers):
        if any(pattern in header for pattern in patterns):
            return index
    return None


def main() -> int:
    if len(sys.argv) != 2:
        raise SystemExit("usage: inspect_price_source.py <xlsx-path>")

    source = Path(sys.argv[1])
    workbook = load_workbook(source, read_only=True, data_only=False, keep_links=True)
    if len(workbook.worksheets) != 1:
        raise SystemExit("SOURCE_WORKBOOK_SHEET_COUNT_UNEXPECTED")

    worksheet = workbook.worksheets[0]
    rows = worksheet.iter_rows(values_only=False)
    first_row = next(rows, None)
    if first_row is None:
        raise SystemExit("SOURCE_WORKBOOK_EMPTY")

    headers = [normalize(cell.value) for cell in first_row]
    if any(not header for header in headers):
        raise SystemExit("SOURCE_WORKBOOK_HEADER_EMPTY")
    if len(set(headers)) != len(headers):
        raise SystemExit("SOURCE_WORKBOOK_HEADER_DUPLICATE")

    pii_markers = ("nome", "cpf", "cnpj", "email", "e-mail", "telefone", "celular", "endereco", "endereço", "corretor", "cliente", "comprador")
    disallowed_markers = pii_markers + ("status", "venda", "vendido", "reserva", "contrato", "pagamento", "parcela", "comissao", "comissão", "repasse", "cobranca", "cobrança")
    pii_headers = [header for header in headers if any(marker in header for marker in pii_markers)]
    disallowed_headers = [header for header in headers if any(marker in header for marker in disallowed_markers)]

    block_index = index_for(headers, ("quadra", "bloco"))
    lot_index = index_for(headers, ("lote",))
    area_index = index_for(headers, ("área", "area", "m²", "m2"))
    price_index = index_for(headers, ("preço", "preco", "valor", "r$"))

    type_counts = [Counter() for _ in headers]
    formula_count = 0
    nonempty_rows = 0
    blank_identifier_rows = 0
    pairs: set[tuple[str, str]] = set()
    duplicate_pairs = 0
    area_numeric_count = 0
    price_numeric_count = 0

    for cells in rows:
        values = [cell.value for cell in cells]
        if not any(value not in (None, "") for value in values):
            continue
        nonempty_rows += 1
        for index, cell in enumerate(cells):
            type_counts[index][classify(cell.value)] += 1
            if isinstance(cell.value, str) and cell.value.startswith("="):
                formula_count += 1

        if block_index is not None and lot_index is not None:
            block = normalize(values[block_index])
            lot = normalize(values[lot_index])
            if not block or not lot:
                blank_identifier_rows += 1
            elif (block, lot) in pairs:
                duplicate_pairs += 1
            else:
                pairs.add((block, lot))
        if area_index is not None and is_number(values[area_index]):
            area_numeric_count += 1
        if price_index is not None and is_number(values[price_index]):
            price_numeric_count += 1

    report = {
        "workbook": {"sheet_count": len(workbook.worksheets), "sheet_name": worksheet.title, "row_count_including_header": worksheet.max_row, "column_count": worksheet.max_column},
        "headers": headers,
        "restricted_headers_detected": sorted(set(disallowed_headers)),
        "personal_data_headers_detected": sorted(set(pii_headers)),
        "formula_count": formula_count,
        "external_links_detected": len(workbook._external_links),
        "data_rows_nonempty": nonempty_rows,
        "column_type_counts": {headers[index]: dict(type_counts[index]) for index in range(len(headers))},
        "structural_reconciliation": {
            "block_column_detected": block_index is not None,
            "lot_column_detected": lot_index is not None,
            "area_column_detected": area_index is not None,
            "price_column_detected": price_index is not None,
            "unique_block_lot_pairs": len(pairs),
            "duplicate_block_lot_pairs": duplicate_pairs,
            "blank_block_or_lot_rows": blank_identifier_rows,
            "numeric_area_rows": area_numeric_count,
            "numeric_price_rows": price_numeric_count,
        },
    }
    print(json.dumps(report, ensure_ascii=False, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
